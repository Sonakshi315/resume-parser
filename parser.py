
import re
from pathlib import Path

import fitz
from docx import Document

SKILLS = [
    "python", "java", "c", "c++", "sql", "mysql",
    "html", "css", "javascript", "react", "node.js",
    "machine learning", "deep learning", "nlp",
    "data analysis", "pandas", "numpy", "matplotlib",
    "scikit-learn", "tensorflow", "pytorch", "excel",
    "power bi", "tableau", "git", "github", "linux",
    "aws", "docker", "flask", "django", "streamlit",
    "rest api", "data structures", "algorithms"
]


def extract_text(file_path):
    path = Path(file_path)

    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    if path.suffix.lower() == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))

        return "\n".join(parts)

    raise ValueError("Only PDF and DOCX files are supported.")


def extract_email(text):
    match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )
    return match.group(0) if match else "Not found"


def extract_phone(text):
    pattern = r"(?<!\w)\+?\d[\d\s().-]{7,}\d(?!\w)"

    for match in re.finditer(pattern, text):
        number = match.group(0).strip()
        digits = re.sub(r"\D", "", number)

        if 10 <= len(digits) <= 13:
            return number

    return "Not found"


def extract_name(text):
    ignored = {
        "resume", "curriculum vitae", "contact",
        "education", "experience", "skills"
    }
    job_titles = {
        "developer", "software", "engineer", "analyst", "manager",
        "specialist", "consultant", "architect", "intern", "student",
        "lead", "designer", "technical", "recruiter", "data"
    }
    skill_tokens = {
        token.lower()
        for skill in SKILLS
        for token in re.split(r"[\s+/]+", skill)
        if token
    }

    for line in text.splitlines()[:12]:
        line = line.strip().rstrip(":")

        if not line or len(line) > 60 or "@" in line:
            continue

        if any(char.isdigit() for char in line):
            continue

        lowered = line.lower()
        if lowered in ignored or "http" in lowered:
            continue

        words = [word.strip(".,;()[]{}") for word in line.split()]

        if not (2 <= len(words) <= 4):
            continue

        if any(word.lower() in job_titles or word.lower() in skill_tokens for word in words):
            continue

        if any(word.lower() in {"mr", "ms", "mrs", "miss", "dr", "prof"} for word in words[:-1]):
            continue

        if not all(re.fullmatch(r"[A-Za-z][A-Za-z'.-]*", word) for word in words):
            continue

        if lowered.startswith("resume ") or lowered.startswith("curriculum vitae"):
            continue

        return line

    return "Not confidently detected"


def extract_skills(text):
    found = []
    normalized = text.lower()

    for skill in SKILLS:
        pattern = (
            r"(?<![a-z0-9+#])"
            + re.escape(skill)
            + r"(?![a-z0-9+#])"
        )

        if re.search(pattern, normalized):
            found.append(skill)

    return sorted(found)


def extract_education(text):
    keywords = [
        "b.tech", "b.e.", "bachelor", "master", "m.tech",
        "mca", "bca", "b.sc", "m.sc", "mba",
        "university", "college", "school", "education"
    ]

    return [
        line.strip()
        for line in text.splitlines()
        if line.strip()
        and any(k in line.lower() for k in keywords)
    ][:15]


def extract_experience(text):
    keywords = [
        "intern", "experience", "employment",
        "worked at", "work history"
    ]

    return [
        line.strip()
        for line in text.splitlines()
        if line.strip()
        and any(k in line.lower() for k in keywords)
    ][:15]


def parse_resume(file_path):
    text = extract_text(file_path)

    if not text.strip():
        raise ValueError(
            "No selectable text found. This PDF may need OCR."
        )

    return {
        "name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "skills": extract_skills(text),
        "education": extract_education(text),
        "experience": extract_experience(text),
        "raw_text": text
    }


def match_job_description(skills, job_description):
    job_skills = set(extract_skills(job_description))
    resume_skills = set(skills)

    matched = sorted(job_skills & resume_skills)
    missing = sorted(job_skills - resume_skills)

    percentage = (
        round(len(matched) / len(job_skills) * 100, 1)
        if job_skills else 0.0
    )

    return {
        "matched_skills": matched,
        "missing_skills": missing,
        "match_percentage": percentage
    }
